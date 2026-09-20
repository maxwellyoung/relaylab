#!/usr/bin/env python3
"""Build the submission DOCX from the repository's report evidence."""

from __future__ import annotations

import re
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "TECHNICAL_REPORT.md"
LAB_EVIDENCE = ROOT / "docs" / "LAB_EVIDENCE.md"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "59636E"
LIGHT_FILL = "F2F4F7"
BORDER = "B7C1CC"
CONTENT_WIDTH_DXA = 9029
TABLE_INDENT_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_font(run, name: str = "Calibri", size: float | None = None, *,
             color: str | None = None, bold: bool | None = None,
             italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top: int = 80, start: int = 120,
                     bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start),
                        ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER)
        borders.append(element)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table columns must total {CONTENT_WIDTH_DXA} DXA")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_properties.append(cannot_split)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb("1F252B")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.left_indent = Inches(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_sections(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header.paragraph_format.space_after = Pt(0)
        run = header.add_run("RELAYLAB  |  COMP713 ASSESSMENT 2  |  MAXWELL YOUNG  |  23213801")
        set_font(run, size=8.5, color=MUTED, bold=True)

        footer = section.footer.paragraphs[0]
        add_page_field(footer)


def add_inline(paragraph, text: str) -> None:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Menlo", size=8.5, color=NAVY)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, italic=True)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def table_widths(column_count: int) -> list[int]:
    patterns = {
        2: [2700, 6660],
        3: [3000, 2520, 3840],
        4: [1900, 2500, 2460, 2500],
    }
    if column_count in patterns:
        return [width * CONTENT_WIDTH_DXA // 9360 for width in patterns[column_count]]
    return [CONTENT_WIDTH_DXA // column_count] * column_count


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    widths = table_widths(len(rows[0]))
    if sum(widths) != CONTENT_WIDTH_DXA:
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = 1.05
            add_inline(paragraph, value.strip())
            for run in paragraph.runs:
                set_font(run, size=9.2, bold=(row_index == 0),
                         color=(NAVY if row_index == 0 else "1F252B"))
            if row_index == 0:
                set_cell_fill(cell, LIGHT_FILL)
                keep_with_next(paragraph)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F6F8FA")
    p_pr.append(shd)
    run = paragraph.add_run("\n".join(lines))
    set_font(run, name="Menlo", size=8.2, color=NAVY)


def load_diagram_font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/SFNS.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def add_architecture_diagram(document: Document) -> None:
    canvas = Image.new("RGB", (1600, 600), "white")
    draw = ImageDraw.Draw(canvas)
    title_font = load_diagram_font(33, bold=True)
    body_font = load_diagram_font(23)
    small_font = load_diagram_font(20)

    boxes = {
        "browser": (70, 125, 390, 290),
        "coordinator": (640, 125, 960, 290),
        "downstream": (1210, 125, 1530, 290),
        "database": (640, 405, 960, 570),
    }
    labels = {
        "browser": ("React client", "Port 5173"),
        "coordinator": ("Coordinator API", "Port 3000"),
        "downstream": ("Downstream service", "Port 3001"),
        "database": ("Relational store", "SQLite or MySQL"),
    }
    for key, box in boxes.items():
        fill = "#F2F4F7" if key != "coordinator" else "#E8F1FA"
        draw.rounded_rectangle(box, radius=18, fill=fill, outline="#7A93AA", width=4)
        left, top, right, bottom = box
        title, subtitle = labels[key]
        title_box = draw.textbbox((0, 0), title, font=title_font)
        subtitle_box = draw.textbbox((0, 0), subtitle, font=body_font)
        draw.text(((left + right - (title_box[2] - title_box[0])) / 2, top + 43),
                  title, fill="#203748", font=title_font)
        draw.text(((left + right - (subtitle_box[2] - subtitle_box[0])) / 2, top + 102),
                  subtitle, fill="#59636E", font=body_font)

    def arrow(start, end, label):
        draw.line([start, end], fill="#2E74B5", width=7)
        x2, y2 = end
        if start[0] == end[0]:
            draw.polygon([(x2, y2), (x2 - 14, y2 - 24), (x2 + 14, y2 - 24)], fill="#2E74B5")
            label_x, label_y = x2 + 28, (start[1] + end[1]) / 2 - 18
        else:
            draw.polygon([(x2, y2), (x2 - 24, y2 - 14), (x2 - 24, y2 + 14)], fill="#2E74B5")
            label_box = draw.textbbox((0, 0), label, font=small_font)
            label_x = (start[0] + end[0] - (label_box[2] - label_box[0])) / 2
            label_y = start[1] - 48
        draw.text((label_x, label_y), label, fill="#2E74B5", font=small_font)

    arrow((390, 207), (640, 207), "JSON over HTTP")
    arrow((960, 207), (1210, 207), "JSON-RPC 2.0")
    arrow((800, 290), (800, 405), "parameterised SQL")

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
        path = Path(handle.name)
    try:
        canvas.save(path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run().add_picture(str(path), width=Inches(6.25))
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_after = Pt(8)
        run = caption.add_run("Figure 1. RelayLab communication and persistence boundaries")
        set_font(run, size=9, color=MUTED, italic=True)
    finally:
        path.unlink(missing_ok=True)


def parse_markdown(document: Document, text: str, *, skip_title: bool = True,
                   heading_shift: int = 0) -> None:
    lines = text.splitlines()
    index = 0
    paragraph_buffer: list[str] = []

    def flush() -> None:
        if paragraph_buffer:
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if skip_title and index == 0 and stripped.startswith("# "):
            index += 1
            while index < len(lines) and (not lines[index].strip() or lines[index].startswith("**")):
                index += 1
            continue

        if stripped.startswith("```"):
            flush()
            block: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            if any("React browser client" in item for item in block):
                add_architecture_diagram(document)
            else:
                add_code_block(document, block)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            flush()
            raw_rows: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                cells = [cell.strip() for cell in candidate.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    raw_rows.append(cells)
                index += 1
            add_table(document, raw_rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush()
            source_level = len(heading.group(1))
            if source_level == 1 and skip_title:
                index += 1
                continue
            level = min(3, max(1, source_level - 1 + heading_shift))
            document.add_paragraph(heading.group(2), style=f"Heading {level}")
            index += 1
            continue

        if stripped.startswith("- "):
            flush()
            bullet_parts = [stripped[2:]]
            index += 1
            while index < len(lines):
                continuation = lines[index].strip()
                if (not continuation or continuation.startswith("- ") or
                        continuation.startswith("#") or continuation.startswith("|") or
                        continuation.startswith("```")):
                    break
                bullet_parts.append(continuation)
                index += 1
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, " ".join(bullet_parts))
            continue

        if not stripped:
            flush()
        else:
            paragraph_buffer.append(stripped)
        index += 1
    flush()


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(110)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("COMP713 DISTRIBUTED AND MOBILE SYSTEMS")
    set_font(run, size=10, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("RelayLab")
    set_font(run, size=32, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("A Distributed API Reliability Workbench")
    set_font(run, size=16, color=DARK_BLUE)

    option = document.add_paragraph()
    option.alignment = WD_ALIGN_PARAGRAPH.CENTER
    option.paragraph_format.space_after = Pt(116)
    run = option.add_run("Assessment 2 - Individual Project, Option A")
    set_font(run, size=11, color=MUTED, italic=True)

    for label, value in (
        ("Student", "Maxwell Young"),
        ("Student ID", "23213801"),
        ("Status date", re.search(r"\*\*Implementation status date:\*\* (.+)", REPORT.read_text(encoding="utf-8")).group(1)),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        label_run = paragraph.add_run(f"{label}: ")
        set_font(label_run, size=10.5, color=MUTED, bold=True)
        value_run = paragraph.add_run(value)
        set_font(value_run, size=10.5, color=MUTED)

    document.add_page_break()


def add_contents(document: Document) -> None:
    document.add_paragraph("Contents", style="Heading 1")
    for item in (
        "1. Project Introduction and Requirements",
        "2. Architecture and Technology Stack",
        "3. Implemented Functionality",
        "4. Communication and Distributed-System Concepts",
        "5. Data Design or Message Design",
        "6. Testing and Evidence",
        "7. Limitations and Possible Improvements",
        "8. Running Instructions",
        "References",
        "Appendix A - Detailed lab evidence",
    ):
        paragraph = document.add_paragraph(style="List Number")
        paragraph.style = document.styles["Normal"]
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(5)
        add_inline(paragraph, item)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(0)
    report_text = REPORT.read_text(encoding="utf-8")
    body = report_text.split("\n## 1.", 1)[1].split("\n## References", 1)[0]
    # The fenced diagram is rendered as an image, so its words are not in the document.
    body = re.sub(r"```.*?```", "", body, flags=re.S)
    report_words = len(re.findall(r"\b[\w'-]+\b", body))
    if report_words > 2000:
        raise ValueError(f"Report exceeds the 2,000-word limit: {report_words}")
    run = note.add_run(
        f"Main body word count (sections 1-8): {report_words:,} words. "
        "The title page, contents, references and appendix are excluded "
        "from the 2,000-word limit."
    )
    set_font(run, size=9.5, color=MUTED, italic=True)
    document.add_page_break()


def build(output: Path) -> None:
    document = Document()
    document.core_properties.author = "Maxwell Young"
    document.core_properties.title = "RelayLab: a distributed API reliability workbench"
    document.core_properties.comments = "COMP713 Assessment 2, Individual Project, Option A"

    configure_styles(document)
    configure_sections(document)
    add_cover(document)
    add_contents(document)
    parse_markdown(document, REPORT.read_text(encoding="utf-8"))
    document.add_page_break()
    document.add_paragraph("Appendix A - Detailed lab evidence", style="Heading 1")
    lab_text = LAB_EVIDENCE.read_text(encoding="utf-8")
    parse_markdown(document, lab_text, skip_title=True, heading_shift=1)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: build-report.py OUTPUT.docx")
    build(Path(sys.argv[1]).resolve())
